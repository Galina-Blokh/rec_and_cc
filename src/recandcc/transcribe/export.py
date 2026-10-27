"""Export a finished session to DOCX and PDF (``transcript.docx`` / ``transcript.pdf``).

Reads ``transcript.json`` (plain text; the bidi isolates are only in md/srt) and writes one
paragraph per segment: ``[hh:mm:ss] Speaker:`` and the text. Hebrew segments are laid out
right to left: DOCX gets a bidi paragraph and an rtl run; the PDF is shaped by HarfBuzz and
reordered with fpdf2's bidi support. The PDF needs a system font that covers Latin, Cyrillic
and Hebrew (Arial, Arial Unicode, DejaVu Sans, ...); without one it raises ExportError. The
libraries are imported lazily, so the app start-up doesn't pay for them. Offline; transcript
text is never logged.
"""

from __future__ import annotations

import contextlib
import functools
import io
import json
import logging
import os
import threading
import time
import unicodedata
from pathlib import Path

from ..models import Segment, is_rtl
from .output import JSON_NAME, fmt_hms, languages_of

log = logging.getLogger(__name__)

DOCX_NAME, PDF_NAME = "transcript.docx", "transcript.pdf"
FORMATS = ("docx", "pdf")
_FONT_CANDIDATES = (  # smaller faces first: Arial Unicode is ~22 MB
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
)


class ExportError(Exception):
    """The session can't be exported (message is safe to show: no transcript text)."""


def rtl_segment(seg: Segment) -> bool:
    """Hebrew by language, or (no language) by the first strong character."""
    return is_rtl(seg.language, seg.text)


def _clean(text: str) -> str:
    """Control characters (they are invalid in DOCX XML) become spaces."""
    return "".join(" " if unicodedata.category(c) == "Cc" else c for c in text).strip()


def load_session(folder: Path) -> tuple[dict, list[Segment]]:
    path = Path(folder) / JSON_NAME
    try:
        data = json.loads(path.read_text("utf-8"))
        segs = [
            Segment(
                float(s["start"]),
                float(s["end"]),
                _clean(str(s["text"])),
                _clean(str(s.get("speaker", ""))),
                str(s.get("language", "")),
            )
            for s in data["segments"]
        ]
    except FileNotFoundError:
        raise ExportError(f"no {JSON_NAME} in {folder}") from None
    except (OSError, ValueError, KeyError, TypeError, AttributeError) as e:
        raise ExportError(f"can't read {path}: {type(e).__name__}") from None
    return data, [s for s in segs if s.text]


def _title(data: dict) -> str:
    started = str(data.get("started_at", ""))[:16].replace("T", " ")
    return f"Transcript {started}".strip()


def _meta(data: dict, segments: list[Segment]) -> str:
    app = data.get("trigger_app") or "manual"
    langs = ", ".join(languages_of(segments)) or "-"
    speakers = ", ".join(dict.fromkeys(s.speaker for s in segments if s.speaker)) or "-"
    return (
        f"Date: {str(data.get('started_at', ''))[:19].replace('T', ' ')} | Trigger: {app} | "
        f"Duration: {fmt_hms(float(data.get('duration') or 0))} | Languages: {langs} | "
        f"Speakers: {speakers}"
    )


_REPLACE_TRIES, _REPLACE_WAIT_S = 20, 0.1  # ~2 s: Word, a viewer or sync may hold the file
_export_lock = threading.Lock()  # one export at a time in this process


def _create_exclusive(tmp: Path) -> int:
    """Create the temp file (0600, O_EXCL). A stale file of the same name must not block it;
    on Windows a scanner may hold it for a moment, so removal and creation are retried."""
    for attempt in range(_REPLACE_TRIES):
        try:
            tmp.unlink(missing_ok=True)
            return os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        except (PermissionError, FileExistsError):
            if attempt + 1 == _REPLACE_TRIES:
                raise
            time.sleep(_REPLACE_WAIT_S)
    raise AssertionError("unreachable")


# POSIX can pin the checked folder with a directory descriptor; Windows can't (residual risk).
_DIR_FD = os.open in os.supports_dir_fd and os.replace in os.supports_dir_fd


def _write(path: Path, data: bytes, expected: Path | None = None) -> Path:
    """Write ``data`` to a temp file and rename it into place. The temp file is created
    owner-only (0600) and written through that same descriptor (no reopen that a scanner
    could block), with a name unique per process and thread (on Windows the output ACLs are
    inherited from the owner-only session folder, #42). Windows may hold the target open, so
    only the rename is retried. Only exception types are kept: library messages can quote the
    transcript.

    ``expected`` is the session folder resolved when the export started. The write is refused
    if the folder now resolves elsewhere (it, or a parent, was swapped for a link or junction
    while the document was rendered); a folder that was already a link at the start is fine.
    On POSIX the checked folder is then pinned by a directory descriptor, so a swap after the
    check can't redirect the write; on Windows a swap in that millisecond remains possible.
    Without ``expected`` any link or junction is refused."""
    parent = path.parent
    if expected is None:
        if parent.is_symlink() or parent.is_junction():
            raise ExportError("The session folder is a link; not writing through it")
    elif parent.resolve() != expected:
        raise ExportError("The session folder moved during the export; not writing through it")
    tmp_name = f".{path.name}.{os.getpid()}.{threading.get_ident()}.tmp"
    if expected is not None and _DIR_FD:
        _write_in_dir(expected, path.name, tmp_name, data)
        return path
    tmp = parent / tmp_name
    try:
        try:
            fd = _create_exclusive(tmp)
            with os.fdopen(fd, "wb") as f:
                f.write(data)
        except Exception as e:
            raise ExportError(f"could not write {path.name} ({type(e).__name__})") from None
        for attempt in range(_REPLACE_TRIES):
            try:
                os.replace(tmp, path)
                return path
            except PermissionError:
                if attempt + 1 < _REPLACE_TRIES:
                    time.sleep(_REPLACE_WAIT_S)
            except OSError as e:
                raise ExportError(f"could not write {path.name} ({type(e).__name__})") from None
        raise ExportError(f"{path.name} may be open in another app; close it and try again")
    finally:
        with contextlib.suppress(OSError):
            tmp.unlink(missing_ok=True)


def _write_in_dir(directory: Path, name: str, tmp_name: str, data: bytes) -> None:
    """POSIX: create, write and rename inside the already-checked folder through one
    directory descriptor (no path lookups after the check)."""
    nofollow = getattr(os, "O_NOFOLLOW", 0)
    try:
        dfd = os.open(directory, os.O_RDONLY | os.O_DIRECTORY | nofollow)
    except OSError as e:
        raise ExportError(f"could not write {name} ({type(e).__name__})") from None
    try:
        with contextlib.suppress(FileNotFoundError):
            os.unlink(tmp_name, dir_fd=dfd)
        try:
            flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL | nofollow
            fd = os.open(tmp_name, flags, 0o600, dir_fd=dfd)
            with os.fdopen(fd, "wb") as f:
                f.write(data)
            os.replace(tmp_name, name, src_dir_fd=dfd, dst_dir_fd=dfd)
        except OSError as e:
            raise ExportError(f"could not write {name} ({type(e).__name__})") from None
    finally:
        with contextlib.suppress(OSError):
            os.unlink(tmp_name, dir_fd=dfd)
        os.close(dfd)


STALE_TEMP_AGE_S = 3600.0


def _sweep_stale_temps(folder: Path) -> None:
    """Remove export temp files (never the transcript's own ``.transcript.md.tmp`` etc.) that
    a crashed export left behind: ``.transcript.docx.<pid>.<tid>.tmp`` / ``.pdf``, untouched
    for an hour (so another process's export in flight is left alone)."""
    # No cross-process lock is needed: only temps untouched for an hour are removed, and an
    # export in flight (here or in another process) writes its temp within seconds.
    now = time.time()
    for name in (DOCX_NAME, PDF_NAME):
        for stale in Path(folder).glob(f".{name}.*.tmp"):
            with contextlib.suppress(OSError):
                if now - stale.stat().st_mtime > STALE_TEMP_AGE_S:
                    stale.unlink()


def _type_only(fn):
    """Library failures become ExportError with the exception type only: their messages can
    quote the transcript."""

    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        try:
            return fn(*args, **kwargs)
        except ExportError:
            raise
        except Exception as e:
            raise ExportError(
                f"{fn.__name__.removeprefix('export_').upper()} failed ({type(e).__name__})"
            ) from None

    return wrapper


@_type_only
def export_docx(folder: Path, loaded=None) -> Path:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    expected = Path(folder).resolve()
    data, segments = loaded or load_session(folder)
    doc = Document()
    doc.core_properties.title = _title(data)
    doc.add_heading(_title(data), level=1)
    doc.add_paragraph(_meta(data, segments))
    if not segments:
        doc.add_paragraph("No speech detected.")
    for seg in segments:
        label = f"[{fmt_hms(seg.start)}] {seg.speaker}:"
        if rtl_segment(seg):
            # the LTR label gets its own paragraph, or the bidi paragraph would mirror it
            lp = doc.add_paragraph()
            lp.add_run(label).bold = True
            lp.paragraph_format.space_after = Pt(0)
            p = doc.add_paragraph()
            p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))  # before w:spacing (schema)
            p.add_run(seg.text)._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
        else:
            p = doc.add_paragraph()
            p.add_run(label + " ").bold = True
            p.add_run(seg.text)
        p.paragraph_format.space_after = Pt(6)
    buf = io.BytesIO()
    doc.save(buf)
    return _write(Path(folder) / DOCX_NAME, buf.getvalue(), expected)


# Latin (with an accent), Cyrillic (with ё), Hebrew letters, niqqud, gershayim and geresh,
# and common punctuation: a font with only a few glyphs of each script is not enough.
_PROBE = "aZé\u0416\u0451\u044f\u05d0\u05ea\u05b7\u05f3\u05f4—…«»“”"


def _covers_scripts(path: Path) -> bool:
    """True if the font has glyphs for Latin, Cyrillic and Hebrew (fontTools comes with fpdf2)."""
    try:
        from fontTools.ttLib import TTFont

        with TTFont(str(path), lazy=True, fontNumber=0) as f:
            cmap = f.getBestCmap() or {}
        return all(ord(c) in cmap for c in _PROBE)
    except Exception:  # noqa: BLE001 - an unreadable font is just not a candidate
        return False


@functools.lru_cache(maxsize=1)
def find_font() -> Path:
    windir = Path(os.environ.get("WINDIR", "C:/Windows"))
    for c in (*map(Path, _FONT_CANDIDATES), windir / "Fonts" / "arial.ttf"):
        if c.is_file() and _covers_scripts(c):
            return c
    raise ExportError(
        "no system font with Latin, Cyrillic and Hebrew found for the PDF (install Arial or "
        "DejaVu Sans); DOCX export still works"
    )


@_type_only
def export_pdf(folder: Path, loaded=None) -> Path:
    from fpdf import FPDF
    from fpdf.enums import WrapMode

    expected = Path(folder).resolve()
    data, segments = loaded or load_session(folder)
    font = find_font()
    pdf = FPDF(format="A4")
    pdf.set_title(_title(data))
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, 18)
    pdf.add_font("body", "", str(font))
    pdf.set_text_shaping(True)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def para(text: str, size: int, rtl: bool = False, gap: float = 2) -> None:
        pdf.set_font("body", size=size)
        pdf.set_text_shaping(True, direction="rtl" if rtl else "ltr")
        pdf.set_x(pdf.l_margin)
        long_token = any(len(t) > 40 for t in text.split())
        pdf.multi_cell(
            width,
            size * 0.5,
            text,
            align="R" if rtl else "L",
            wrapmode=WrapMode.CHAR if long_token else WrapMode.WORD,
            new_x="LMARGIN",
            new_y="NEXT",
        )
        pdf.ln(gap)

    para(_title(data), 16, gap=3)
    para(_meta(data, segments), 9, gap=5)
    if not segments:
        para("No speech detected.", 11)
    for seg in segments:
        label = f"[{fmt_hms(seg.start)}] {seg.speaker}:"
        if rtl_segment(seg):
            # the LTR label on its own line keeps it intact above the right-to-left body
            para(label, 11, gap=0)
            para(seg.text, 11, rtl=True, gap=3)
        else:
            para(f"{label} {seg.text}", 11, gap=3)
    return _write(Path(folder) / PDF_NAME, bytes(pdf.output()), expected)


def export_session(folder: Path, fmt: str = "both", skipped: list[str] | None = None) -> list[Path]:
    """Write the requested files next to the transcript; returns their paths. Both files are
    built from one read of transcript.json. With ``both``, a PDF that can't be made (no
    font) doesn't lose the DOCX: the reason is appended to ``skipped`` if one is given."""
    if fmt not in (*FORMATS, "both"):
        raise ExportError(f"unknown format {fmt!r}")
    with _export_lock:
        _sweep_stale_temps(folder)
        loaded = load_session(folder)
        out = []
        if fmt in ("docx", "both"):
            out.append(export_docx(folder, loaded))
        if fmt in ("pdf", "both"):
            try:
                out.append(export_pdf(folder, loaded))
            except Exception as e:
                err = (
                    e
                    if isinstance(e, ExportError)
                    else ExportError(f"PDF failed ({type(e).__name__})")
                )
                if fmt == "pdf":
                    raise err from None
                if skipped is not None:
                    skipped.append(str(err))
    log.info("exported session %s (%s)", Path(folder).name, fmt)
    return out
