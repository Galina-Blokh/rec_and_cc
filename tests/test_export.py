"""DOCX / PDF export of a session transcript (synthetic transcript.json only).

Stories: US-25.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn

from recandcc import cli
from recandcc.models import Segment, SessionInfo
from recandcc.transcribe import export

SEGS = [
    {"start": 0.0, "end": 2.0, "text": "Hello team", "speaker": "Me", "language": "en"},
    {"start": 3.0, "end": 5.0, "text": "Привет, давайте", "speaker": "Others", "language": "ru"},
    {"start": 6.0, "end": 8.0, "text": "שלום לכולם", "speaker": "Speaker 1", "language": "he"},
]


@pytest.fixture
def session(tmp_path):
    d = tmp_path / "2026-10-26" / "140312_zoom-us"
    d.mkdir(parents=True)
    data = {
        "session_id": "x",
        "started_at": "2026-10-26T14:03:12",
        "trigger_app": "zoom-us",
        "duration": 9.0,
        "segments": SEGS,
    }
    (d / "transcript.json").write_text(json.dumps(data, ensure_ascii=False), "utf-8")
    return d


def have_font() -> bool:
    try:
        export.find_font()
    except export.ExportError:
        return False
    return True


def test_docx_has_text_speakers_timestamps_and_rtl(session):
    path = export.export_docx(session)
    assert path == session / "transcript.docx"
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert "Transcript 2026-10-26 14:03" in texts[0]
    assert "[00:00:00] Me: Hello team" in texts
    assert "[00:00:03] Others: Привет, давайте" in texts
    # a Hebrew segment: the LTR label has its own paragraph, the text a bidi one
    assert texts[-2:] == ["[00:00:06] Speaker 1:", "שלום לכולם"]
    bidi = [
        p._p.pPr is not None and p._p.pPr.find(qn("w:bidi")) is not None for p in doc.paragraphs
    ]
    assert bidi[-1] and not any(bidi[:-1])
    assert doc.paragraphs[-2].runs[0].bold


def test_docx_bidi_comes_before_spacing_in_schema_order(session):
    ppr = Document(str(export.export_docx(session))).paragraphs[-1]._p.pPr
    tags = [c.tag for c in ppr]
    assert tags.index(qn("w:bidi")) < tags.index(qn("w:spacing"))


def test_control_characters_are_cleaned(tmp_path):
    (tmp_path / "transcript.json").write_text(
        json.dumps(
            {"segments": [{"start": 0, "end": 1, "text": "a\x00b\x0bc", "speaker": "M\x01e"}]}
        ),
        "utf-8",
    )
    _data, segs = export.load_session(tmp_path)
    assert (segs[0].text, segs[0].speaker) == ("a b c", "M e")
    Document(str(export.export_docx(tmp_path)))  # valid XML


def test_temp_file_is_private_from_the_start(session, monkeypatch):
    if sys.platform == "win32":
        pytest.skip("POSIX modes")
    modes, real = [], export.os.replace

    def spy(src, dst):
        modes.append(os.stat(src).st_mode & 0o777)
        return real(src, dst)

    monkeypatch.setattr(export.os, "replace", spy)
    export._write(session / "x.bin", b"data")
    assert modes == [0o600] and (session / "x.bin").read_bytes() == b"data"


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_pdf_text_reads_back_in_all_three_languages(session):
    from pypdf import PdfReader

    text = "\n".join(pg.extract_text() for pg in PdfReader(str(export.export_pdf(session))).pages)
    assert "Transcript 2026-10-26 14:03" in text
    assert "[00:00:00] Me: Hello team" in text and "Привет, давайте" in text
    he = "שלום לכולם"
    assert he in text or he[::-1] in text  # logical or visual order, depending on the extractor


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_a_very_long_unbroken_token_does_not_break_the_pdf(tmp_path):
    (tmp_path / "transcript.json").write_text(
        json.dumps({"segments": [{"start": 0, "end": 1, "text": "x" * 400, "speaker": "Me"}]}),
        "utf-8",
    )
    assert export.export_pdf(tmp_path).read_bytes().startswith(b"%PDF")


def test_font_without_hebrew_is_not_picked(monkeypatch, tmp_path):
    f = tmp_path / "a.ttf"
    f.write_bytes(b"not a font")
    export.find_font.cache_clear()
    monkeypatch.setattr(export, "_FONT_CANDIDATES", (str(f),))
    monkeypatch.setenv("WINDIR", str(tmp_path / "none"))
    with pytest.raises(export.ExportError):
        export.find_font()


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_pdf_is_written(session):
    path = export.export_pdf(session)
    raw = path.read_bytes()
    assert raw.startswith(b"%PDF") and b"%%EOF" in raw[-32:]


def test_export_session_formats_and_private_mode(session):
    out = export.export_session(session, "docx")
    assert [p.name for p in out] == ["transcript.docx"]
    if sys.platform != "win32":
        assert (session / "transcript.docx").stat().st_mode & 0o077 == 0
    assert not list(session.glob(".*.tmp"))
    with pytest.raises(export.ExportError):
        export.export_session(session, "odt")


def test_clean_maps_surrogates_and_noncharacters_so_docx_still_exports(session):
    assert export._clean("a\ud83db\uffffc\x01d") == "a\ufffdb\ufffdc d"
    data = json.loads((session / "transcript.json").read_text("utf-8"))
    data["segments"][0]["text"] = "x\uffff y"
    (session / "transcript.json").write_text(json.dumps(data), "utf-8")
    assert [p.name for p in export.export_session(session, "docx")] == ["transcript.docx"]


def test_missing_or_broken_transcript(tmp_path):
    with pytest.raises(export.ExportError):
        export.export_docx(tmp_path)
    (tmp_path / "transcript.json").write_text("{nope", "utf-8")
    with pytest.raises(export.ExportError):
        export.export_docx(tmp_path)


def test_rtl_detection_by_text_when_language_unknown():
    assert export.rtl_segment(Segment(0, 1, "שלום", "Me", ""))
    assert not export.rtl_segment(Segment(0, 1, "hello", "Me", ""))


def test_cli_export(session, capsys):
    assert cli.main(["export", str(session), "--format", "docx"]) == 0
    assert (session / "transcript.docx").exists()
    assert str(session / "transcript.docx") in capsys.readouterr().out
    assert cli.main(["export", str(session / "nope"), "--format", "docx"]) == 1


def test_cli_rejects_unknown_format(session):
    with pytest.raises(SystemExit):
        cli.main(["export", str(session), "--format", "odt"])


def test_importing_the_app_does_not_load_export_libraries():
    code = (
        "import sys, recandcc, recandcc.cli, recandcc.transcribe.export;"
        "sys.exit(any(m in sys.modules for m in ('docx', 'fpdf')))"
    )
    assert subprocess.run([sys.executable, "-c", code], check=False, timeout=60).returncode == 0


def _controller(output_dir=None):
    import queue
    import threading

    from recandcc.ui.common import Controller

    ctrl = Controller.__new__(Controller)  # only the export wiring, no runtime
    ctrl.events, ctrl._export_lock = queue.Queue(), threading.Lock()
    export_cfg = type("E", (), {"template": ""})()
    ctrl.cfg = type("C", (), {"output_dir": output_dir, "export": export_cfg})()
    return ctrl


@pytest.mark.parametrize("fmt", ["docx", "pdf", "both"])
def test_controller_export_posts_a_notice(session, monkeypatch, fmt):
    seen = []

    def fake(folder, f, skipped, template=None):
        seen.append(f)
        (folder / "transcript.docx").write_text("x")
        return [folder / "transcript.docx"]

    monkeypatch.setattr(export, "export_session", fake)
    ctrl = _controller(session.parent.parent)
    ctrl.export_async(session / "transcript.md", fmt)
    kind, title, msg = ctrl.events.get(timeout=10)
    assert (kind, title) == ("notice", "Export") and msg.startswith("Saved transcript.docx (")
    assert seen == [fmt]


def test_controller_releases_the_lock_when_the_thread_cannot_start(session, monkeypatch):
    import threading

    def no_start(self):
        raise RuntimeError("can't start new thread")

    monkeypatch.setattr(threading.Thread, "start", no_start)
    ctrl = _controller(session.parent.parent)
    with pytest.raises(RuntimeError):
        ctrl.export_async(session / "transcript.md")
    assert ctrl._export_lock.acquire(blocking=False)


def test_controller_export_logs_only_the_error_type(session, monkeypatch, caplog):
    def boom(*_a, **_kw):
        raise RuntimeError("secret words from the call")

    monkeypatch.setattr(export, "export_session", boom)
    ctrl = _controller(session.parent.parent)
    caplog.set_level("DEBUG")
    ctrl.export_async(session / "transcript.md")
    assert ctrl.events.get(timeout=10)[2].startswith("Export failed")
    assert "RuntimeError" in caplog.text and "secret words" not in caplog.text


def test_controller_ignores_a_second_export_while_one_runs(session):
    ctrl = _controller(session.parent.parent)
    ctrl._export_lock.acquire()
    ctrl.export_async(session / "transcript.md")
    assert ctrl.events.empty()


def test_library_errors_become_export_errors_without_text(session, monkeypatch):
    import docx

    def boom(self, _path):
        raise ValueError("Привет, секрет")

    monkeypatch.setattr(docx.document.Document, "save", boom)
    with pytest.raises(export.ExportError) as e:
        export.export_docx(session)
    assert "ValueError" in str(e.value) and "секрет" not in str(e.value)
    assert not list(session.glob(".*.tmp"))


def test_missing_font_keeps_the_docx_when_both(session, monkeypatch):
    def nofont():
        raise export.ExportError("no system font")

    monkeypatch.setattr(export, "find_font", nofont)
    skipped: list[str] = []
    out = export.export_session(session, "both", skipped)
    assert [p.name for p in out] == ["transcript.docx"] and skipped == ["no system font"]
    with pytest.raises(export.ExportError):
        export.export_session(session, "pdf")


def test_both_formats_share_one_snapshot(session, monkeypatch):
    calls = []
    real = export.load_session
    monkeypatch.setattr(export, "load_session", lambda f, *a: calls.append(f) or real(f, *a))
    export.export_session(session, "both")
    assert len(calls) == 1


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_pdf_lays_out_each_segment_with_its_direction(session, monkeypatch):
    from fpdf import FPDF

    seen = []
    real = FPDF.multi_cell

    def spy(self, w, h=None, text="", *a, **kw):
        seen.append((text, kw.get("align")))
        return real(self, w, h, text, *a, **kw)

    monkeypatch.setattr(FPDF, "multi_cell", spy)
    export.export_pdf(session)
    assert ("[00:00:00] Me: Hello team", "L") in seen
    assert ("[00:00:03] Others: Привет, давайте", "L") in seen
    assert ("[00:00:06] Speaker 1:", "L") in seen and ("שלום לכולם", "R") in seen


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_pdf_file_has_title_and_one_page(session):
    import re

    raw = export.export_pdf(session).read_bytes()
    assert b"Transcript 2026-10-26 14:03" in raw  # the document title (metadata)
    assert len(re.findall(rb"/Type /Page\b", raw)) == 1


def test_concurrent_exports_of_one_session_are_serialised(session):
    from concurrent.futures import ThreadPoolExecutor

    with ThreadPoolExecutor(4) as pool:
        results = list(pool.map(lambda _: export.export_session(session, "docx"), range(4)))
    assert all(r == [session / "transcript.docx"] for r in results)
    Document(str(session / "transcript.docx"))  # a whole, readable file
    assert not list(session.glob(".*.tmp"))


def test_a_pdf_failure_other_than_the_font_keeps_the_docx(session, monkeypatch):
    def boom(folder, loaded=None, summary=None, **kw):
        raise RuntimeError("Привет, секрет")

    monkeypatch.setattr(export, "export_pdf", boom)
    skipped: list[str] = []
    out = export.export_session(session, "both", skipped)
    assert [p.name for p in out] == ["transcript.docx"]
    assert skipped == ["PDF failed (RuntimeError)"]


def test_replace_is_retried_while_the_file_is_held_open(session, monkeypatch):
    monkeypatch.setattr(export, "_DIR_FD", False)  # the Windows (path-based) write
    real, calls = export.os.replace, []

    def flaky(src, dst):
        calls.append(1)
        if len(calls) < 3:
            raise PermissionError("in use")
        return real(src, dst)

    monkeypatch.setattr(export.os, "replace", flaky)
    monkeypatch.setattr(export.time, "sleep", lambda _s: None)
    export.export_docx(session)
    assert len(calls) == 3
    Document(str(session / "transcript.docx"))
    assert not list(session.glob(".*.tmp"))


def test_replace_that_never_succeeds_is_an_export_error(session, monkeypatch):
    def locked(src, dst):
        raise PermissionError("in use")

    monkeypatch.setattr(export.os, "replace", locked)
    monkeypatch.setattr(export.time, "sleep", lambda _s: None)
    with pytest.raises(export.ExportError):
        export.export_docx(session)
    assert not list(session.glob(".*.tmp"))


def test_latest_finished_session_skips_live_ones_and_ignores_mtime(tmp_path):
    from recandcc.ui.common import latest_finished_session

    def sess(day, name, done):
        d = tmp_path / day / name
        d.mkdir(parents=True)
        (d / "transcript.md").write_text("x")
        if done:
            (d / "transcript.json").write_text("{}")
        return d / "transcript.md"

    assert latest_finished_session(tmp_path) is None
    for i in range(12):  # many newer live sessions
        sess("2026-10-27", f"{i:02d}0000_zoom", False)
    new = sess("2026-10-26", "140312_zoom", True)
    old = sess("2026-10-25", "100000_meet", True)
    os.utime(old, (9e9, 9e9))  # re-transcribed just now: newest mtime, oldest session
    assert latest_finished_session(tmp_path) == new
    (new.parent / SessionInfo.DISCARD_MARKER).write_text("")  # discarded: the next one counts
    assert latest_finished_session(tmp_path) == old
    assert latest_finished_session(tmp_path / "missing") is None


def test_controller_export_of_nothing_finished_posts_a_notice(tmp_path):
    ctrl = _controller(tmp_path)
    ctrl.export_async()
    assert ctrl.events.get(timeout=10) == ("notice", "Export", "No finished transcripts yet")


def test_a_held_open_file_gets_a_clear_notice(session, monkeypatch):
    monkeypatch.setattr(export, "_DIR_FD", False)  # the Windows (path-based) write

    def locked(src, dst):
        raise PermissionError("in use")

    monkeypatch.setattr(export.os, "replace", locked)
    monkeypatch.setattr(export.time, "sleep", lambda _s: None)
    with pytest.raises(export.ExportError, match="open in another app"):
        export.export_docx(session)


def test_export_refuses_a_session_linked_from_outside_the_folder(tmp_path):
    if sys.platform == "win32":
        pytest.skip("symlinks need privileges")
    out, elsewhere = tmp_path / "out", tmp_path / "elsewhere"
    (out / "2099-12-31").mkdir(parents=True)
    elsewhere.mkdir()
    (out / "2099-12-31" / "000000_x").symlink_to(elsewhere)
    ctrl = _controller(out)
    ctrl.export_async(out / "2099-12-31" / "000000_x" / "transcript.md")
    assert "outside" in ctrl.events.get(timeout=10)[2]
    assert not list(elsewhere.iterdir())


def test_only_old_export_temps_are_swept(session):
    old = session / ".transcript.docx.99999.1.tmp"
    fresh = session / ".transcript.pdf.99999.1.tmp"  # another process, in flight
    others = [session / f".transcript.{e}.tmp" for e in ("md", "json", "srt")]
    for f in (old, fresh, *others):
        f.write_text("x")
    os.utime(old, (1, 1))
    export.export_session(session, "docx")
    assert not old.exists() and fresh.exists() and all(f.exists() for f in others)


def test_a_write_error_is_not_reported_as_a_file_in_use(session, monkeypatch):
    real = export.os.open

    def denied(path, *a, **kw):
        if str(path).endswith(".tmp"):
            raise PermissionError("read-only folder")
        return real(path, *a, **kw)

    monkeypatch.setattr(export.os, "open", denied)
    monkeypatch.setattr(export.time, "sleep", lambda _s: None)
    with pytest.raises(export.ExportError, match=r"could not write.*PermissionError"):
        export.export_docx(session)


def test_a_failing_cleanup_does_not_hide_the_export_error(session, monkeypatch):
    monkeypatch.setattr(export, "_DIR_FD", False)  # the Windows (path-based) write
    from pathlib import Path

    def locked(src, dst):
        raise PermissionError("in use")

    calls = []

    def no_unlink(self, missing_ok=False):
        calls.append(1)
        if len(calls) > 1:  # the pre-create removal works; the final cleanup is blocked
            raise PermissionError("held by antivirus")

    monkeypatch.setattr(export.os, "replace", locked)
    monkeypatch.setattr(export.time, "sleep", lambda _s: None)
    monkeypatch.setattr(Path, "unlink", no_unlink)
    with pytest.raises(export.ExportError, match="open in another app"):
        export.export_docx(session)


def test_a_long_meeting_exports(tmp_path):
    segs = [
        {
            "start": i,
            "end": i + 1,
            "text": f"line {i} שלום привет",
            "speaker": "Me",
            "language": "he",
        }
        for i in range(300)
    ]
    (tmp_path / "transcript.json").write_text(
        json.dumps({"segments": segs}, ensure_ascii=False), "utf-8"
    )
    assert len(Document(str(export.export_docx(tmp_path))).paragraphs) > 600
    if have_font():
        assert export.export_pdf(tmp_path).read_bytes().startswith(b"%PDF")


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_a_wrapped_rtl_paragraph_keeps_its_line_order():
    from fpdf import FPDF

    words = [f"מילה{chr(0x5D0 + i % 26)}{chr(0x5D0 + i // 26)}" for i in range(60)]
    pdf = FPDF(format="A4")
    pdf.add_font("body", "", str(export.find_font()))
    pdf.set_text_shaping(True, direction="rtl")
    pdf.add_page()
    pdf.set_font("body", size=11)
    lines = pdf.multi_cell(150, 6, " ".join(words), align="R", dry_run=True, output="LINES")
    assert len(lines) > 1  # it wrapped
    assert " ".join(lines).split() == words  # logical order kept across the wrap points


def test_a_session_folder_swapped_for_a_link_is_not_written_through(tmp_path):
    if sys.platform == "win32":
        pytest.skip("symlinks need privileges")
    real, link = tmp_path / "real", tmp_path / "link"
    real.mkdir()
    link.symlink_to(real)
    with pytest.raises(export.ExportError, match="link"):
        export._write(link / "x.bin", b"data")
    assert not list(real.iterdir())


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_pdf_safe_text_placeholder_bidi_and_scripts():
    font = export.find_font()
    cmap = export._font_cmap(font)
    if 0xE000 in cmap:
        pytest.skip("font covers the private-use probe")
    out = export.pdf_safe_text("ab‫c⁦d؜", font)
    assert out[0] == "a" and out[1] in "?�" and out[2:] == "bcd"
    kept = "Привет שלום\n"
    assert export.pdf_safe_text(kept, font) == kept
    # LRM/RLM survive only where the font has a glyph; never a placeholder
    marks = export.pdf_safe_text("x‏y‎z", font)
    assert marks.replace("‏", "").replace("‎", "") == "xyz"
    assert ("‏" in marks) == (0x200F in cmap)


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_pdf_safe_text_collapses_emoji_sequences():
    font = export.find_font()
    cmap = export._font_cmap(font)
    if any(o in cmap for o in (0x1F468, 0x1F469, 0x1F466, 0x1F44D, 0x1F3FD)):
        pytest.skip("font covers emoji")
    family = "\U0001f468‍\U0001f469‍\U0001f466"
    thumbs = "\U0001f44d\U0001f3fd"
    out = export.pdf_safe_text(f"{family} {thumbs}", font)
    assert out.count("?") + out.count("�") == 2 and " " in out


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_pdf_with_emoji_and_bidi_marks_reads_back(tmp_path):
    from pypdf import PdfReader

    text = "hi ‏ Привет שלום there"
    seg = {"start": 0, "end": 1, "text": text, "speaker": "Me"}
    (tmp_path / "transcript.json").write_text(json.dumps({"segments": [seg]}), "utf-8")
    pdf = export.export_pdf(tmp_path)
    got = "\n".join(pg.extract_text() for pg in PdfReader(str(pdf)).pages)
    he = "שלום"
    assert "Привет" in got and (he in got or he[::-1] in got)
    assert "?" in got or "�" in got
    assert "\ue000" not in got


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_pdf_safe_text_keeps_a_placeholder_after_zwj_and_normalises_controls():
    font = export.find_font()
    cmap = export._font_cmap(font)
    if 0x1F468 in cmap or 0x4E2D in cmap or 0xE000 in cmap:
        pytest.skip("font covers the probes")
    # a lone ZWJ after an uncovered emoji must not swallow the next, non-emoji character
    out = export.pdf_safe_text("\U0001f468‍中", font)
    assert [c in "?�" for c in out] == [True, True]
    assert export.pdf_safe_text("a\r\nb\rc\x00d\x07e", font) == "a\nb\ncde"


def test_a_parent_swapped_for_a_link_during_rendering_is_refused(tmp_path):
    if sys.platform == "win32":
        pytest.skip("symlinks need privileges")
    day, other = tmp_path / "day", tmp_path / "other"
    (day / "s").mkdir(parents=True)
    (other / "s").mkdir(parents=True)
    with export._pinned(day / "s") as pin:
        day.rename(tmp_path / "day-old")
        day.symlink_to(other)  # an ancestor became a link; the session folder itself is not one
        with pytest.raises(export.ExportError, match="moved"):
            export._write(day / "s" / "x.bin", b"data", pin)
    assert not list((other / "s").iterdir())
    assert not list((tmp_path / "day-old" / "s").iterdir())


def test_a_junction_session_folder_is_refused_without_expected(tmp_path, monkeypatch):
    # runs on every platform: the junction check itself is faked
    (tmp_path / "s").mkdir()
    monkeypatch.setattr(export.Path, "is_junction", lambda self: True)
    with pytest.raises(export.ExportError, match="link"):
        export._write(tmp_path / "s" / "x.bin", b"data")
    assert not list((tmp_path / "s").iterdir())


def test_write_without_expected_into_a_plain_folder_is_unchanged(tmp_path):
    out = export._write(tmp_path / "x.bin", b"data")
    assert out.read_bytes() == b"data"
    assert [p.name for p in tmp_path.iterdir()] == ["x.bin"]  # no temp left behind


def test_a_session_folder_that_was_already_a_link_at_the_start_is_exported(tmp_path):
    # e.g. a redirected or synced folder: linked before the export began, so it's legitimate
    if sys.platform == "win32":
        pytest.skip("symlinks need privileges")
    real, link = tmp_path / "real", tmp_path / "link"
    real.mkdir()
    link.symlink_to(real)
    with export._pinned(link) as pin:
        out = export._write(link / "x.bin", b"data", pin)
    assert (real / "x.bin").read_bytes() == b"data" and out == link / "x.bin"
    assert [p.name for p in real.iterdir()] == ["x.bin"]


def test_the_dir_fd_write_path_is_used_on_posix():
    # the descriptor-pinned write must really run on macOS and Linux, not silently fall back
    if sys.platform == "win32":
        pytest.skip("no dir_fd on Windows: documented residual risk")
    assert export._DIR_FD


posix_only = pytest.mark.skipif(sys.platform == "win32", reason="symlinks and dir_fd: POSIX")


def _swap_for_link(folder, target, keep_as):
    folder.rename(keep_as)
    folder.symlink_to(target)


@posix_only
def test_an_ancestor_swapped_between_resolve_and_stat_is_caught(tmp_path, monkeypatch):
    # the folder is resolved, then opened (and fstat'ed): a parent swapped in that gap makes
    # the open follow the new link, so the identity would be the swapped folder's. Resolving
    # again after the open catches it.
    day, other = tmp_path / "day", tmp_path / "other"
    (day / "s").mkdir(parents=True)
    (other / "s").mkdir(parents=True)
    real_open = export.os.open

    def swap_then_open(p, *a, **kw):
        if "dir_fd" not in kw and not (tmp_path / "day-old").exists():
            _swap_for_link(day, other, tmp_path / "day-old")
        return real_open(p, *a, **kw)

    monkeypatch.setattr(export.os, "open", swap_then_open)
    with pytest.raises(export.ExportError, match="moved"):
        export._pin_folder(day / "s")


@posix_only
def test_a_swap_after_the_pin_cannot_redirect_the_write(tmp_path):
    # even if the path check were passed, the bytes go through the pinned descriptor
    (tmp_path / "s").mkdir()
    (tmp_path / "evil").mkdir()
    with export._pinned(tmp_path / "s") as pin:
        _swap_for_link(tmp_path / "s", tmp_path / "evil", tmp_path / "s-moved")
        assert pin.dfd is not None
        export._write_in_dir(pin.dfd, "x.bin", ".x.tmp", b"data")
    assert not list((tmp_path / "evil").iterdir())
    assert (tmp_path / "s-moved" / "x.bin").read_bytes() == b"data"


@posix_only
def test_a_swap_between_the_read_and_the_write_is_refused(session, tmp_path, monkeypatch):
    evil = tmp_path / "evil"
    evil.mkdir()
    real = export.load_summary

    def read_then_swap(*a, **kw):
        out = real(*a, **kw)
        _swap_for_link(session, evil, session.with_name("moved"))
        return out

    monkeypatch.setattr(export, "load_summary", read_then_swap)
    with pytest.raises(export.ExportError, match="moved"):
        export.export_session(session, "docx")
    assert not list(evil.iterdir())
    assert not list(session.with_name("moved").glob("transcript.docx"))


@posix_only
@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_a_swap_between_docx_and_pdf_does_not_redirect_the_pdf(session, tmp_path, monkeypatch):
    evil = tmp_path / "evil"
    evil.mkdir()
    real_docx = export.export_docx

    def docx_then_swap(*a, **kw):
        out = real_docx(*a, **kw)
        _swap_for_link(session, evil, session.with_name("moved"))
        return out

    monkeypatch.setattr(export, "export_docx", docx_then_swap)
    skipped: list[str] = []
    out = export.export_session(session, "both", skipped)
    assert [p.name for p in out] == ["transcript.docx"]
    assert skipped and "moved" in skipped[0]
    assert not list(evil.iterdir())


@posix_only
def test_a_write_error_in_the_pinned_folder_is_an_export_error(session, monkeypatch):
    def failing_replace(src, dst, *, src_dir_fd=None, dst_dir_fd=None):
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(export.os, "replace", failing_replace)
    with pytest.raises(export.ExportError, match=r"could not write transcript\.docx \(OSError\)"):
        export.export_docx(session)
    assert not list(session.glob(".*.tmp"))  # the temp is removed through the descriptor


@posix_only
def test_a_read_only_pinned_folder_is_an_export_error(session):
    if os.geteuid() == 0:
        pytest.skip("root ignores directory permissions")
    session.chmod(0o500)
    try:
        with pytest.raises(export.ExportError, match="could not write.*PermissionError"):
            export.export_docx(session)
    finally:
        session.chmod(0o700)


@posix_only
def test_a_folder_linked_from_the_start_keeps_its_summary_in_the_export(session, tmp_path):
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    link = tmp_path / "linked"
    link.symlink_to(session)
    export.export_session(link, "docx")
    paras = [p.text for p in Document(str(session / "transcript.docx")).paragraphs]
    assert "Action items" in paras


SUMMARY_MD = (
    "# Summary\n\nWe agreed to ship on Friday.\n\n## Action items\n\n"
    "- Anna sends the release notes\n\n_Generated on this computer by m; it can be wrong._\n"
)
SUMMARY_HE = (
    "# סיכום\n\n\u2067סיכמנו לשחרר ביום שישי\u2069\n\n## משימות\n\n- \u2067לשלוח הערות\u2069\n"
)


def test_docx_without_summary_is_unchanged_in_structure(session):
    plain = [p.text for p in Document(str(export.export_docx(session))).paragraphs]
    assert "Transcript" not in plain and "Summary" not in plain
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    with_summary = [p.text for p in Document(str(export.export_docx(session))).paragraphs]
    assert with_summary[:2] == plain[:2]  # title and meta line
    assert with_summary[-len(plain[2:]) :] == plain[2:]  # transcript paragraphs untouched


def test_docx_puts_summary_and_action_items_before_the_transcript(session):
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    texts = [p.text for p in Document(str(export.export_docx(session))).paragraphs]
    for needle in ("Summary", "Action items", "Anna sends the release notes", "Transcript"):
        assert needle in texts
    assert texts.index("Action items") < texts.index("Transcript")
    assert texts.index("Transcript") < texts.index("[00:00:00] Me: Hello team")


def test_docx_hebrew_summary_is_bidi_without_isolates(session):
    (session / "summary.md").write_text(SUMMARY_HE, "utf-8")
    doc = Document(str(export.export_docx(session)))
    by_text = {p.text: p for p in doc.paragraphs}
    p = by_text["סיכמנו לשחרר ביום שישי"]
    assert p._p.pPr.find(qn("w:bidi")) is not None
    assert "לשלוח הערות" in by_text and "\u2067" not in "".join(by_text)


@pytest.mark.skipif(not have_font(), reason="no system font with Latin/Cyrillic/Hebrew")
def test_pdf_has_summary_first_and_hebrew_survives(session):
    from pypdf import PdfReader

    def pdf_text() -> str:
        return "\n".join(
            pg.extract_text() for pg in PdfReader(str(export.export_pdf(session))).pages
        )

    assert "Action items" not in pdf_text()
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    text = pdf_text()
    assert "Action items" in text and "Anna sends the release notes" in text
    assert text.index("Action items") < text.index("Hello team")
    (session / "summary.md").write_text(SUMMARY_HE, "utf-8")
    text = pdf_text()
    assert "\u05dc\u05e9\u05dc\u05d5\u05d7 \u05d4\u05e2\u05e8\u05d5\u05ea" in text
    assert text.index("\u05de\u05e9\u05d9\u05de\u05d5\u05ea") < text.index("Hello team")


def test_an_unreadable_summary_is_ignored(session):
    (session / "summary.md").write_bytes(b"\xff\xfe not utf-8 \x80")
    assert export.load_summary(session) == []
    assert export.export_docx(session).is_file()


def test_summary_over_the_cap_is_skipped(session, caplog):
    (session / "summary.md").write_text("# Summary\n" + "x" * export._MAX_SUMMARY_BYTES, "utf-8")
    with caplog.at_level("DEBUG", logger="recandcc.transcribe.export"):
        assert export.load_summary(session) == []
    (session / "summary.md").write_text("# Summary\nbody\n", "utf-8")
    assert export.load_summary(session)


def test_a_malformed_summary_logs_only_the_error_type(session, caplog):
    (session / "summary.md").write_bytes(b"\x80secret words")
    with caplog.at_level("DEBUG", logger="recandcc.transcribe.export"):
        assert export.load_summary(session) == []
    assert "UnicodeDecodeError" in caplog.text and "secret" not in caplog.text


@pytest.mark.skipif(sys.platform == "win32", reason="symlinks need privileges on Windows")
def test_a_linked_summary_or_folder_is_refused(session, tmp_path):
    outside = tmp_path / "outside.md"
    outside.write_text("# Summary\nsecret\n", "utf-8")
    (session / "summary.md").symlink_to(outside)
    assert export.load_summary(session) == []
    link = tmp_path / "linked"
    link.symlink_to(session)
    (session / "summary.md").unlink()
    (session / "summary.md").write_text("# Summary\nbody\n", "utf-8")
    # a folder that was a link from the start is legitimate (redirected or synced folders)
    assert export.load_summary(link) == [("h1", "Summary"), ("p", "body")]


def test_a_bom_does_not_hide_the_heading(session):
    (session / "summary.md").write_bytes(b"\xef\xbb\xbf# Summary\n\nBody\n")
    assert export.load_summary(session)[0] == ("h1", "Summary")


def test_markdown_variants(session):
    (session / "summary.md").write_text(
        "-\n* star\n+ plus\n1. one\n2) two\n### Sub\n_a_b_c_\ntext\n", "utf-8"
    )
    assert export.load_summary(session) == [
        ("bullet", "star"),
        ("bullet", "plus"),
        ("num", "1. one"),
        ("num", "2. two"),
        ("h3", "Sub"),
        ("note", "a_b_c"),
        ("p", "text"),
    ]


def test_docx_summary_bullet_and_note_directions(session):
    (session / "summary.md").write_text(
        "- Anna לשלוח הערות\n- Anna sends notes\n_נוצר במחשב_\n",
        "utf-8",
    )
    paras = Document(str(export.export_docx(session))).paragraphs
    he = next(p for p in paras if p.text.startswith("Anna ל"))  # mixed, Hebrew majority
    en = next(p for p in paras if p.text == "Anna sends notes")
    note = next(p for p in paras if p.text == "נוצר במחשב")
    assert he._p.pPr.find(qn("w:bidi")) is not None
    assert en._p.pPr is None or en._p.pPr.find(qn("w:bidi")) is None
    rpr = note.runs[0]._r.rPr
    assert rpr.find(qn("w:i")) is not None and rpr.find(qn("w:iCs")) is not None


@pytest.mark.skipif(not have_font(), reason="no system font with Latin/Cyrillic/Hebrew")
def test_pdf_hebrew_bullet_sits_at_the_right_margin(session):
    import re

    from pypdf import PdfReader
    from pypdf.generic import ContentStream

    (session / "summary.md").write_text(
        "# s\n\n- \u05dc\u05e9\u05dc\u05d5\u05d7 \u05d4\u05e2\u05e8\u05d5\u05ea\n", "utf-8"
    )
    page = PdfReader(str(export.export_pdf(session))).pages[0]
    font = next(iter(page["/Resources"]["/Font"].values())).get_object()
    cmap = font["/ToUnicode"].get_object().get_data().decode("latin-1")
    code = {int(c, 16): chr(int(u, 16)) for c, u in re.findall(r"<(\w{4})>\s*<(\w{4})>", cmap)}
    glyphs: list[tuple[float, str]] = []  # (x, char), one per glyph the PDF draws
    x = 0.0
    for operands, op in ContentStream(page.get_contents(), page.pdf).operations:
        if op == b"Td":
            x = float(operands[0])
        elif op == b"Tm":
            x = float(operands[4])
        elif op == b"Tj" and operands[0]:
            glyphs.append(
                (x, code.get(int.from_bytes(operands[0].encode("latin-1")[:2], "big"), ""))
            )
    row = [(gx, ch) for gx, ch in glyphs if ch == "\u2022" or "\u05d0" <= ch <= "\u05ea"]
    assert any(ch == "\u2022" for _, ch in row)
    assert max(row)[1] == "\u2022"  # the rightmost glyph on the line is the bullet


def test_a_pending_or_older_summary_is_not_exported(session):
    import os

    from recandcc.transcribe.summary import PENDING_MARKER

    summary = session / "summary.md"
    summary.write_text(SUMMARY_MD, "utf-8")
    assert export.load_summary(session)
    (session / PENDING_MARKER).touch()  # being redone, e.g. after a speaker rename
    assert export.load_summary(session) == []
    (session / PENDING_MARKER).unlink()
    t = (session / "transcript.json").stat().st_mtime
    os.utime(summary, (t - 10, t - 10))  # the transcript was redone after the summary
    assert export.load_summary(session) == []
    assert "Action items" not in [
        p.text for p in Document(str(export.export_docx(session))).paragraphs
    ]


@pytest.mark.skipif(not have_font(), reason="no system font with Latin/Cyrillic/Hebrew")
def test_pdf_export_survives_a_corrupt_summary(session):
    (session / "summary.md").write_bytes(b"\xff\xfe\x00 broken \x80")
    assert export.export_pdf(session).read_bytes().startswith(b"%PDF")


def test_docx_bidi_is_in_schema_order(session):
    (session / "summary.md").write_text(SUMMARY_HE, "utf-8")
    doc = Document(str(export.export_docx(session)))
    order = ["pStyle", "numPr", "bidi", "spacing", "ind", "jc"]
    checked = 0
    for p in doc.paragraphs:
        ppr = p._p.pPr
        if ppr is None or ppr.find(qn("w:bidi")) is None:
            continue
        tags = [c.tag.split("}")[1] for c in ppr]
        assert tags == sorted(tags, key=order.index)
        checked += 1
    assert checked >= 3
    bullet = next(
        p
        for p in doc.paragraphs
        if p.text == "\u05dc\u05e9\u05dc\u05d5\u05d7 \u05d4\u05e2\u05e8\u05d5\u05ea"
    )
    tags = [c.tag.split("}")[1] for c in bullet._p.pPr]
    assert tags[:2] == ["pStyle", "numPr"] or tags[0] == "pStyle"
    assert tags.index("bidi") > tags.index("pStyle")


def test_numbered_items_keep_their_numbers(session):
    (session / "summary.md").write_text("# s\n\n1. first\n2) second\n", "utf-8")
    texts = [p.text for p in Document(str(export.export_docx(session))).paragraphs]
    assert "1. first" in texts and "2. second" in texts


def test_numbered_hebrew_item_is_bidi(session):
    (session / "summary.md").write_text("# s\n\n1. \u05dc\u05e9\u05dc\u05d5\u05d7\n", "utf-8")
    p = next(
        p for p in Document(str(export.export_docx(session))).paragraphs if p.text.startswith("1.")
    )
    assert p._p.pPr.find(qn("w:bidi")) is not None


def test_both_formats_make_the_same_summary_decision(session, monkeypatch, caplog):
    import os

    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    real = export.export_docx

    def docx_then_retranscribe(*a, **kw):
        out = real(*a, **kw)
        t = (session / "transcript.json").stat().st_mtime + 100
        os.utime(session / "transcript.json", (t, t))  # transcript redone mid-export
        return out

    monkeypatch.setattr(export, "export_docx", docx_then_retranscribe)
    seen = []
    real_pdf = export.export_pdf
    monkeypatch.setattr(
        export,
        "export_pdf",
        lambda f, loaded=None, summary=None, **kw: (
            seen.append(summary) or real_pdf(f, loaded, summary, **kw)
        ),
    )
    export.export_session(session)
    assert seen and seen[0]  # the PDF got the same, non-empty summary as the DOCX
    assert any(
        p.text == "Action items" for p in Document(str(session / "transcript.docx")).paragraphs
    )


def test_a_stale_summary_is_logged_with_the_reason_only(session, caplog):
    import logging
    import os

    summary = session / "summary.md"
    summary.write_text(SUMMARY_MD, "utf-8")
    t = (session / "transcript.json").stat().st_mtime
    os.utime(summary, (t - 10, t - 10))
    with caplog.at_level(logging.DEBUG, logger="recandcc.transcribe.export"):
        assert export.load_summary(session) == []
    assert "older than the transcript" in caplog.text
    assert "Action items" not in caplog.text


def test_a_year_is_not_a_numbered_item(session):
    (session / "summary.md").write_text("2026. The budget was approved\n12. real item\n", "utf-8")
    assert export.load_summary(session) == [
        ("p", "2026. The budget was approved"),
        ("num", "12. real item"),
    ]


@pytest.mark.skipif(not have_font(), reason="no Latin+Cyrillic+Hebrew system font")
def test_pdf_summary_text_also_gets_a_placeholder(session):
    from pypdf import PdfReader

    if 0xE000 in export._font_cmap(export.find_font()):
        pytest.skip("font covers the private-use probe")
    (session / "summary.md").write_text("# Summary\n\nShip on Friday  ok.\n", "utf-8")
    pdf = export.export_pdf(session)
    got = "\n".join(pg.extract_text() for pg in PdfReader(str(pdf)).pages)
    assert "" not in got and ("Friday ? ok" in got or "Friday � ok" in got)


def _summary_texts(session):
    return [t for _, t in export.load_summary(session)]


def _rewrite_transcript(session, edit):
    path = session / "transcript.json"
    data = json.loads(path.read_text("utf-8"))
    edit(data["segments"])
    path.write_text(json.dumps(data, ensure_ascii=False), "utf-8")
    os.utime(path, (9e9, 9e9))  # a rewrite is always newer than the summary


def _with_fingerprint(session):
    from recandcc.transcribe.summary import SOURCE_NAME, transcript_fingerprint

    segs = json.loads((session / "transcript.json").read_text("utf-8"))["segments"]
    (session / SOURCE_NAME).write_text(transcript_fingerprint(segs), "utf-8")


def test_summary_survives_a_speaker_rename_rewrite(session):
    """Stories: US-25."""
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    _with_fingerprint(session)

    def rename(segs):
        for s in segs:
            s["speaker"] = "Anna"  # name-speaker / call-app names relabel in place

    _rewrite_transcript(session, rename)
    assert _summary_texts(session)


def test_summary_is_dropped_when_the_text_changed(session):
    """Stories: US-25."""
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    _with_fingerprint(session)
    _rewrite_transcript(session, lambda segs: segs[0].update(text="something else entirely"))
    assert _summary_texts(session) == []


def test_summary_without_a_fingerprint_keeps_the_mtime_rule(session):
    """Stories: US-25."""
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    os.utime(session / "summary.md", (9e9, 9e9))
    assert _summary_texts(session)  # newer than the transcript
    os.utime(session / "summary.md", (1, 1))
    assert _summary_texts(session) == []


def test_summarize_session_records_the_fingerprint(session):
    """Stories: US-25."""
    from recandcc.config import Config
    from recandcc.transcribe import summary

    gen = lambda messages, max_tokens: "- do it"  # noqa: E731
    summary.summarize_session(session, Config(), gen)

    def rename(segs):
        segs[0]["speaker"] = "Zed"

    _rewrite_transcript(session, rename)
    assert _summary_texts(session)


def test_a_null_segments_list_is_a_stale_fingerprint_not_a_crash(session):
    """Stories: US-25. ``"segments": null`` makes the fingerprint raise TypeError."""
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    _with_fingerprint(session)
    _rewrite_transcript_data(session, lambda d: d.update(segments=None))
    os.utime(session / "summary.md", (9e9 + 1, 9e9 + 1))
    assert _summary_texts(session)  # falls back to the mtime rule (summary is newer)


def _rewrite_transcript_data(session, edit):
    path = session / "transcript.json"
    data = json.loads(path.read_text("utf-8"))
    edit(data)
    path.write_text(json.dumps(data, ensure_ascii=False), "utf-8")
    os.utime(path, (9e9, 9e9))


@pytest.mark.parametrize("garbage", [b"\xff\xfe\x00 not utf8", b"", b"deadbeef"])
def test_an_unreadable_summary_source_falls_back_to_the_mtime_rule(session, garbage):
    """Stories: US-25."""
    from recandcc.transcribe.summary import SOURCE_NAME

    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    (session / SOURCE_NAME).write_bytes(garbage)
    if garbage == b"deadbeef":  # readable but wrong: the text "changed"
        assert _summary_texts(session) == []
        return
    os.utime(session / "summary.md", (9e9, 9e9))
    assert _summary_texts(session)
    os.utime(session / "summary.md", (1, 1))
    assert _summary_texts(session) == []


def test_a_timing_only_change_makes_the_summary_stale(session):
    """Stories: US-25."""
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    _with_fingerprint(session)
    _rewrite_transcript(session, lambda segs: segs[1].update(start=3.5, end=5.5))
    assert _summary_texts(session) == []


def test_summary_survives_a_rename_through_write_outputs(session):
    """Stories: US-25. The real rewrite path (what name-speaker and the menu use)."""
    from dataclasses import replace
    from datetime import datetime

    from recandcc.models import TriggerKind
    from recandcc.transcribe import output

    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    _with_fingerprint(session)
    info = SessionInfo(
        id="x",
        directory=session,
        started_at=datetime(2026, 10, 26, 14, 3, 12),
        trigger_kind=TriggerKind.CALL,
        trigger_app="zoom-us",
    )
    segs = [replace(Segment(**s), speaker="Anna") for s in SEGS]
    output.write_outputs(info, segs, 9.0, "test-model")
    assert (
        json.loads((session / "transcript.json").read_text("utf-8"))["segments"][0]["speaker"]
        == "Anna"
    )
    os.utime(session / "transcript.json", (9e9, 9e9))  # newer than the summary
    assert _summary_texts(session)


def test_the_pending_marker_wins_over_a_matching_fingerprint(session):
    """Stories: US-25."""
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    _with_fingerprint(session)
    assert _summary_texts(session)
    (session / ".summary-pending").touch()
    assert _summary_texts(session) == []


def test_a_failed_sidecar_write_leaves_no_old_fingerprint(session, monkeypatch):
    """Stories: US-25. The old fingerprint must not drop the fresh summary."""
    from recandcc.config import Config
    from recandcc.transcribe import summary

    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    _with_fingerprint(session)
    _rewrite_transcript(session, lambda segs: segs[0].update(text="new words"))
    real = summary.atomic_write

    def fail_sidecar(path, text):
        if path.name == summary.SOURCE_NAME:
            raise OSError("disk full")
        return real(path, text)

    monkeypatch.setattr(summary, "atomic_write", fail_sidecar)
    summary.summarize_session(session, Config(), lambda messages, max_tokens: "- do it")
    assert not (session / summary.SOURCE_NAME).exists()
    os.utime(session / "summary.md", (9e9 + 1, 9e9 + 1))  # fresh: newer than the transcript
    assert _summary_texts(session)  # mtime rule
