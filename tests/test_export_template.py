"""Custom export templates (``export.template``): synthetic transcript.json only.

Stories: US-25.
"""

from __future__ import annotations

import json
import logging
from pathlib import Path

import pytest
from docx import Document

from recandcc.config import Config, load_config
from recandcc.transcribe import export, export_template

SEGS = [
    {"start": 0.0, "end": 2.0, "text": "Hello team", "speaker": "Me", "language": "en"},
    {"start": 3.0, "end": 5.0, "text": "Привет", "speaker": "Others", "language": "ru"},
]
SUMMARY_MD = (
    "# Summary\n\nWe agreed to ship on Friday.\n\n## Action items\n\n"
    "- Anna sends the release notes\n\n_Generated on this computer by m; it can be wrong._\n"
)
META = (
    "Date: 2026-10-26 14:03:12 | Trigger: zoom-us | Duration: 00:00:09 | "
    "Languages: en, ru | Speakers: Me, Others"
)
PLAIN = [
    "Transcript 2026-10-26 14:03",
    META,
    "[00:00:00] Me: Hello team",
    "[00:00:03] Others: Привет",
]
NOTE = "Generated on this computer by m; it can be wrong."
WITH_SUMMARY = [
    *PLAIN[:2],
    "Summary",
    "We agreed to ship on Friday.",
    "Action items",
    "Anna sends the release notes",
    NOTE,
    "Transcript",
    *PLAIN[2:],
]


@pytest.fixture
def session(tmp_path):
    d = tmp_path / "s"
    d.mkdir()
    data = {
        "started_at": "2026-10-26T14:03:12",
        "trigger_app": "zoom-us",
        "duration": 9.0,
        "segments": SEGS,
    }
    (d / "transcript.json").write_text(json.dumps(data, ensure_ascii=False), "utf-8")
    return d


def texts(session, template=None) -> list[str]:
    export.export_session(session, "docx", template=template)
    return [p.text for p in Document(str(session / "transcript.docx")).paragraphs]


def write_tpl(tmp_path: Path, body: str) -> Path:
    path = tmp_path / "tpl.toml"
    path.write_text(body, "utf-8")
    return path


def test_default_layout_is_unchanged(session):
    assert texts(session) == PLAIN
    assert texts(session, template="") == PLAIN
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    assert texts(session) == WITH_SUMMARY
    assert texts(session, template=None) == WITH_SUMMARY


def test_title_header_footer_and_placeholders(session, tmp_path):
    tpl = write_tpl(
        tmp_path,
        'title = "Meeting {date} at {time} ({duration}) {nope}"\n'
        'header = "Via {app}; {title}"\nfooter = "{speakers} / {languages} {a.b} {0}"\n',
    )
    out = texts(session, str(tpl))
    assert out[0] == "Meeting 2026-10-26 at 14:03 (00:00:09) {nope}"  # unknown stays literal
    assert out[2] == "Via zoom-us; Transcript 2026-10-26 14:03"
    assert out[-1] == "Me, Others / en, ru {a.b} {0}"  # no attribute access, no indexes
    assert out[3:-1] == PLAIN[2:]
    assert Document(str(session / "transcript.docx")).core_properties.title == out[0]


def test_section_order_and_custom_headings(session, tmp_path):
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    tpl = write_tpl(
        tmp_path,
        'sections = ["transcript", "action_items", "summary"]\n'
        '[headings]\nsummary = "Overview"\naction_items = "To do"\ntranscript = "Full text"\n',
    )
    assert texts(session, str(tpl)) == [
        *PLAIN[:2],
        "Full text",
        *PLAIN[2:],
        "To do",
        "Anna sends the release notes",
        "Overview",
        "We agreed to ship on Friday.",
        NOTE,
    ]


def test_sections_can_be_left_out(session, tmp_path):
    (session / "summary.md").write_text(SUMMARY_MD, "utf-8")
    tpl = write_tpl(tmp_path, 'sections = ["summary"]\n')
    assert texts(session, str(tpl)) == [*PLAIN[:2], "Summary", "We agreed to ship on Friday.", NOTE]
    tpl = write_tpl(tmp_path, "sections = []\n")
    assert texts(session, str(tpl)) == PLAIN[:2]


@pytest.mark.parametrize(
    "body",
    ["title = 5", 'sections = ["nope"]', 'sections = "summary"', "[headings]\nx = 1", "= = ="],
)
def test_a_bad_template_warns_once_and_falls_back(session, tmp_path, caplog, body):
    tpl = write_tpl(tmp_path, body + "\n# secret words from the call\n")
    with caplog.at_level(logging.DEBUG):
        assert texts(session, str(tpl)) == PLAIN
    warnings = [r for r in caplog.records if r.levelno == logging.WARNING]
    assert len(warnings) == 1 and "secret words" not in caplog.text


def test_a_missing_template_warns_once_and_falls_back(session, tmp_path, caplog):
    with caplog.at_level(logging.WARNING):
        assert texts(session, str(tmp_path / "gone.toml")) == PLAIN
    assert len([r for r in caplog.records if r.levelno == logging.WARNING]) == 1


def test_template_paths_use_pathlib_and_the_config_folder(session, tmp_path, monkeypatch):
    write_tpl(tmp_path, 'title = "Rel {date}"\n')
    monkeypatch.setattr(export_template, "CONFIG_DIR", tmp_path)
    assert texts(session, "tpl.toml")[0] == "Rel 2026-10-26"  # relative: from the config folder
    assert texts(session, tmp_path / "tpl.toml")[0] == "Rel 2026-10-26"
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    assert texts(session, "~/tpl.toml")[0] == "Rel 2026-10-26"


def test_oversize_template_is_ignored(tmp_path):
    tpl = write_tpl(tmp_path, 'title = "' + "x" * 70_000 + '"\n')
    assert export_template.load_template(str(tpl)) is None


def test_pdf_uses_the_template(session, tmp_path):
    try:
        export.find_font()
    except export.ExportError:
        pytest.skip("no system font")
    tpl = write_tpl(tmp_path, 'title = "Custom {date}"\nfooter = "end"\n')
    path = export.export_session(session, "pdf", template=str(tpl))[0]
    assert path.read_bytes().startswith(b"%PDF")


def test_config_key_defaults_and_loads(tmp_path):
    assert Config().export.template == ""
    cfg_path = tmp_path / "config.toml"
    cfg_path.write_text('[export]\ntemplate = "~/t.toml"\n', "utf-8")
    assert load_config(cfg_path).export.template == "~/t.toml"
    cfg_path.write_text("[export]\ntemplate = 3\n", "utf-8")
    assert load_config(cfg_path).export.template == ""
